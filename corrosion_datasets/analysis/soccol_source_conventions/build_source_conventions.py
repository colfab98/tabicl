"""Build a source-by-source convention registry for the Soccol workbook.

Run from any directory with:
    python build_source_conventions.py

Requires openpyxl. The output is descriptive: it does not modify the raw workbook.
"""

from __future__ import annotations

import csv
import json
import math
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[2]
DATASET_DIR = ROOT / "corrosion_datasets" / "datasets" / "soccol_pitting_potential"
WORKBOOK = DATASET_DIR / "raw" / "Pitting_potential_dataset_2024.xlsx"
METADATA = HERE / "source_reference_metadata.csv"
OUTPUT = HERE / "source_conventions.csv"
SUMMARY = HERE / "source_conventions_summary.json"

ELEMENTS = ["C", "N", "Si", "P", "S", "Ti", "V", "Cr", "Mn", "Ni", "Nb", "Mo"]
MAJOR_ELEMENTS = ["Cr", "Mn", "Ni", "Mo"]
CONDITION_FIELDS = [
    "Prep_grinding_grit", "Prep_Ra_micron", "Prep_medium", "Prep_pH",
    "Prep_redox", "Prep_time", "CP_time", "CP_aeration", "CP_agitation",
    "CP_temp", "CP_pH", "CP_Cl", "CP_Br", "CP_OH", "CP_SO4", "CP_CO3",
    "CP_NO3", "CP_PO4", "CP_MoO4", "CP_CrO4", "CP_ion_other",
    "CP_anions_info", "Test_area_cm2", "scan_rate",
]

COLLECTION_PAPER = "https://doi.org/10.1016/j.electacta.2024.145355"
COLLECTION_MANUSCRIPT = "../../datasets/soccol_pitting_potential/raw/Soccol_updated_PREN_author_manuscript.pdf"

# These overrides are intentionally small. Everything else is derived from the
# workbook, bibliography, and the collection-level paper. They capture sources
# for which the default label "Fe-based stainless steel" would be misleading.
MATERIAL_OVERRIDES: dict[str, dict[str, str]] = {
    "1968Bond": {
        "material_family": "Fe-based austenitic stainless and high-Ni Fe-Ni-Cr alloys",
        "stainless_scope": "yes_with_high_nickel_variants",
        "balance_element": "Fe_expected",
        "balance_reconstruction": "unsafe_on_5_Mo-only_rows; approximate_Fe_balance_elsewhere",
        "material_note": "The paper calls the materials austenitic stainless steels. Five workbook rows contain only Mo composition, so their Fe remainder cannot be reconstructed from this table.",
        "evidence_level": "source_title_and_workbook_compositions",
    },
    "1968Horvath": {
        "material_family": "mixed Ni, Cr-Ni, Cr-Fe model alloys and related stainless steels",
        "stainless_scope": "mixed",
        "balance_element": "row_specific",
        "balance_reconstruction": "do_not_apply_global_Fe_balance",
        "material_note": "The source explicitly spans Ni, Cr-Ni, Cr-Fe and stainless alloys; the workbook includes a 15Cr-59.939Ni row among 15Cr-13Ni-Mo rows.",
        "evidence_level": "source_title_abstract_and_workbook_compositions",
        "evidence_url": "https://doi.org/10.1149/1.2411433",
    },
    "1977Sugimoto": {
        "material_family": "Fe-Cr and Fe-Cr-Ni-Mo experimental model alloys",
        "stainless_scope": "yes_model_alloys",
        "balance_element": "Fe_nominal_balance",
        "balance_reconstruction": "approximate_only; Cr is largest in 60Cr and 70Cr rows",
        "material_note": "Fe is the nominal balance, but it is not the principal element in the 60 and 70 wt% Cr model alloys.",
        "evidence_level": "source_title_and_workbook_compositions",
    },
    "1983Bandy": {
        "material_family": "Fe-based high-alloy experimental stainless steels",
        "stainless_scope": "yes_model_alloys",
        "balance_element": "Fe_nominal_balance",
        "balance_reconstruction": "approximate_only; Fe need not be the largest component",
        "material_note": "Designed Cr-Ni-Mn-Mo-N alloys include rows whose approximate Fe remainder is similar to or below Ni.",
        "evidence_level": "source_title_and_workbook_compositions",
    },
    "1986Jargelius": {
        "material_family": "Fe-based high-N austenitic experimental stainless steels",
        "stainless_scope": "yes_model_alloys",
        "balance_element": "Fe",
        "balance_reconstruction": "approximate_only",
        "material_note": "Later literature cites this conference work as nitrogen-alloying research on austenitic stainless steels; workbook compositions leave approximately 48-53 wt% Fe balance.",
        "evidence_level": "secondary_bibliographic_identification_and_workbook_compositions",
    },
    "1988Roberge": {
        "material_family": "mixed AISI 304, Fe-Ni-Cr Alloy 800, and Ni-based Alloy 600",
        "stainless_scope": "mixed",
        "balance_element": "row_specific_Fe_or_Ni",
        "balance_reconstruction": "split_by_composition_before_balance_reconstruction",
        "material_note": "The six rows are two measurements each for 304, Alloy 800 and Alloy 600.",
        "evidence_level": "source_abstract_and_workbook_compositions",
        "evidence_url": "https://doi.org/10.5006/1.3583937",
    },
    "1991Cortest": {
        "material_family": "mixed Fe-based 1.4438/317L stainless and Ni-based G-3 alloy",
        "stainless_scope": "mixed",
        "balance_element": "Fe_for_1.4438_and_317L; Ni_for_G3",
        "balance_reconstruction": "split_by_alloy_designation",
        "material_note": "G3 rows (22.25Cr-55Ni-7Mo in the workbook) are not Fe-balance stainless steel.",
        "evidence_level": "workbook_designations_and_compositions",
    },
    "1994Carroll": {
        "material_family": "mixed Fe-based 316L/302 stainless and 20Cr-80Ni model alloy",
        "stainless_scope": "mixed",
        "balance_element": "Fe_for_316L_and_302; none_for_20Cr80Ni",
        "balance_reconstruction": "split_rows_where_Cr=20_and_Ni=80",
        "material_note": "Ten unlabeled rows are exactly 20Cr-80Ni; the other rows are 316L or AISI 302.",
        "evidence_level": "source_full_text_metadata_and_workbook_compositions",
        "evidence_url": "https://doi.org/10.1016/0010-938X(94)90061-2",
    },
    "1994Malik": {
        "material_family": "mixed conventional stainless and high-alloy Fe-Ni-Cr grades",
        "stainless_scope": "yes_with_borderline_high_alloy_grades",
        "balance_element": "Fe_expected_by_grade; verify_20Cb3_taxonomy_if_filtering",
        "balance_reconstruction": "approximate_only_due_to_unrecorded_Cu_and_other_grade_elements",
        "material_note": "Includes 20Cb3/Alloy 20 and other highly alloyed grades; a strict stainless-only filter needs an explicit taxonomy rule.",
        "evidence_level": "source_title_and_workbook_designations",
    },
    "1995Malik": {
        "material_family": "mixed conventional stainless and high-alloy Fe-Ni-Cr grades",
        "stainless_scope": "yes_with_borderline_high_alloy_grades",
        "balance_element": "Fe_expected_by_grade; verify_N08020_taxonomy_if_filtering",
        "balance_reconstruction": "approximate_only_due_to_unrecorded_Cu_Ta_and_other_grade_elements",
        "material_note": "Includes UNS N08020/20Cb3 and other highly alloyed grades; some stated Nb+Ta values cannot be separated in the workbook schema.",
        "evidence_level": "source_title_and_workbook_designations",
    },
    "1997Stellwag": {
        "material_family": "Fe-Ni-Cr Alloy 800",
        "stainless_scope": "borderline_related_alloy_not_conventional_stainless",
        "balance_element": "Fe",
        "balance_reconstruction": "approximate_Fe_balance_about_45_wt_pct",
        "material_note": "Alloy 800 is an Fe-Ni-Cr corrosion/heat-resistant alloy and should be flagged separately in a strict stainless-steel task.",
        "evidence_level": "source_title_designation_and_workbook_composition",
    },
    "2000Russell": {
        "material_family": "Fe-based N-Mo stainless experimental alloys",
        "stainless_scope": "yes_model_alloys",
        "balance_element": "Fe",
        "balance_reconstruction": "approximate_after_removing_misfiled_Ti_values",
        "material_note": "The source describes 22Cr-5Ni-5Mn-(1-5)Mo-(0-0.7)N-balance Fe alloys.",
        "evidence_level": "source_abstract_and_formula_verification",
        "evidence_url": "https://doi.org/10.5006/1.3290360",
    },
    "2006Muwila": {
        "material_family": "Fe-based low-Ni austenitic experimental stainless steels",
        "stainless_scope": "yes_model_alloys",
        "balance_element": "Fe",
        "balance_reconstruction": "approximate_only",
        "material_note": "The dissertation reports Hercules-base experimental alloys varying Mn, Mo and N; preliminary button compositions were off target, so measured-versus-target provenance matters.",
        "evidence_level": "institutional_thesis_abstract_and_workbook_compositions",
        "evidence_url": "https://wiredspace.wits.ac.za/items/78b6192d-9bd7-49fb-a478-9f13de00cb81",
    },
    "2006Shin": {
        "material_family": "Fe-based 22Cr-5.5Ni-3.22Mo duplex-like stainless steel",
        "stainless_scope": "yes",
        "balance_element": "Fe",
        "balance_reconstruction": "approximate_only",
        "material_note": "The original title was not recovered; the constant workbook chemistry is a conventional Fe-balance 22Cr-5Ni-3Mo duplex composition.",
        "evidence_level": "workbook_composition_inference",
    },
    "2007Saithala": {
        "material_family": "Fe-based duplex and superduplex stainless steels",
        "stainless_scope": "yes",
        "balance_element": "Fe",
        "balance_reconstruction": "approximate_only",
        "material_note": "The conference source covers Zeron 100, 2205 and Ferralium Alloy 255 duplex grades.",
        "evidence_level": "institutional_conference_record_and_workbook_compositions",
        "evidence_url": "https://shura.shu.ac.uk/view/types/conference%3D5Fitem/2007.date.html",
    },
    "2009Wong": {
        "material_family": "Ni-based Ni-Cr-Mo ternary model alloys",
        "stainless_scope": "no",
        "balance_element": "none; Ni_Cr_Mo_are_explicit_and_sum_about_100",
        "balance_reconstruction": "do_not_add_Fe",
        "material_note": "All 92 rows are Ni-Cr-Mo alloys; recorded Ni+Cr+Mo sums to 99.99-100.01 wt%.",
        "evidence_level": "doctoral_thesis_record_and_workbook_compositions",
        "evidence_url": "https://etd.ohiolink.edu/acprod/odb_etd/r/etd/search/search-results?clear=1001&p1001_keyword=metastable%2520pitting&request=KEYWORD_SEARCH",
    },
    "2015Blackwood": {
        "material_family": "Fe-based 304L and 316L stainless steels",
        "stainless_scope": "yes",
        "balance_element": "Fe",
        "balance_reconstruction": "approximate_only",
        "material_note": "The source abstract identifies 304L and 316L; the DOI stored upstream is truncated and should end in .253.",
        "evidence_level": "source_abstract_and_workbook_compositions",
        "evidence_url": "https://doi.org/10.14773/cst.2015.14.6.253",
    },
}

# Factors explicitly named by the source/designation but absent or misleading
# in the workbook composition schema. These are not inferred blank=zero cases.
SOURCE_ISSUES: dict[str, tuple[str, str]] = {
    "1968Bond": ("five_rows_have_only_Mo_composition", "Do not reconstruct Fe for the five Mo-only rows without the source table."),
    "1968Horvath": ("mixed_base_alloys", "Split Cr-Ni/Ni-rich rows from Fe-Cr/Fe-based stainless rows."),
    "1969Boehni": ("Re_alloying_variable_has_no_workbook_column", "Keep source flag or recover Re from the paper before composition-only modeling."),
    "1969Lizlovs": ("Ti_named_in_designations_but_Ti_cells_are_blank", "Recover Ti from designation/source; zero-filled Ti would be wrong for the Ti-bearing rows."),
    "1974Machin": ("Cu_named_in_high_alloy_designation_but_no_Cu_column", "Treat Fe remainder as approximate and retain designation/source indicator."),
    "1983Bui": ("W_alloying_variable_has_no_workbook_column", "Composition-only rows omit the studied W variation; recover W or flag the source."),
    "1988Roberge": ("mixed_base_alloys", "Split AISI 304, Alloy 800 and Alloy 600 rows before any balance calculation."),
    "1990Ives": ("Mo_implantation_not_equivalent_to_bulk_Mo_column", "Retain source/treatment identity; do not interpret bulk Mo alone as implanted dose."),
    "1991Cortest": ("mixed_base_alloys_and_G3_elements_outside_schema", "Split G3 from Fe-based grades; do not reconstruct G3 as Fe balance."),
    "1992Sabot": ("N_and_Ne_ion_implantation_not_fully_represented_as_bulk_composition", "Retain treatment/source identity or recover implant dose/species."),
    "1994Carroll": ("ten_20Cr80Ni_rows_mixed_with_stainless", "Split rows with Cr=20 and Ni=80 from the Fe-based subset."),
    "1994Malik": ("Cu_and_some_high_alloy_elements_absent_from_schema", "Use grade/source indicators; Fe remainder is approximate."),
    "1995Malik": ("Cu_and_Ta_absent;_Nb+Ta_not_separable", "Use grade/source indicators; Fe remainder is approximate."),
    "1998Ahn": ("W_alloying_variable_has_no_workbook_column", "Recover W before using this source to learn W effects."),
    "2000DiSchino": ("430TiNb_designation_has_no_Nb_value", "Do not interpret blank Nb as verified zero for the 430TiNb row."),
    "2000Russell": ("Ti_column_contains_PRE=Cr+3.3Mo+20N", "Set these 24 Ti cells to missing/drop the field; they are not Ti composition."),
    "2008Pohjanne": ("all_67_rows_have_no_numeric_composition", "Use alloy designation or recover grade chemistry; zero fill erases material differences."),
    "2014Zhou": ("only_S_is_recorded_for_316F_and_316L", "Use grade/source information or recover bulk chemistry; zero fill erases Cr/Ni/Mo."),
    "2018Shi": ("Y_alloying_variable_has_no_workbook_column", "Recover Y content or retain a source/material indicator."),
    "2020Sander": ("all_3_rows_have_no_numeric_composition", "Use 316L designation or recover grade chemistry."),
}

MIXED_FAMILY = {"1968Horvath", "1988Roberge", "1991Cortest", "1994Carroll"}
NON_FE = {"2009Wong"}
REPAIR_REQUIRED = {"2000Russell"}
BORDERLINE = {"1997Stellwag", "1994Malik", "1995Malik"}


def clean(value: Any) -> Any:
    if isinstance(value, str):
        return value.replace("\ufeff", "").strip()
    return value


def finite_number(value: Any) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)


def normalized_tuple(row: dict[str, Any], fields: list[str]) -> tuple[Any, ...]:
    out = []
    for field in fields:
        value = row.get(field)
        if finite_number(value):
            out.append(round(float(value), 12))
        elif value is None:
            out.append(None)
        else:
            out.append(str(value).strip())
    return tuple(out)


def semicolon(values: list[str]) -> str:
    return ";".join(values)


def load_rows() -> list[dict[str, Any]]:
    wb = load_workbook(WORKBOOK, read_only=True, data_only=True)
    ws = wb["pitting_potentials"]
    iterator = ws.iter_rows(values_only=True)
    headers = [clean(x) for x in next(iterator)]
    return [dict(zip(headers, (clean(x) for x in row))) for row in iterator]


def load_metadata() -> dict[str, dict[str, str]]:
    with METADATA.open(newline="", encoding="utf-8-sig") as handle:
        return {row["source"]: row for row in csv.DictReader(handle)}


def main() -> None:
    rows = load_rows()
    metadata = load_metadata()
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        grouped[str(row["source"])].append(row)

    records: list[dict[str, Any]] = []
    for source in sorted(grouped):
        group = grouped[source]
        meta = metadata[source]
        n = len(group)

        missing = {element: sum(row.get(element) is None for row in group) for element in ELEMENTS}
        partial = [e for e in ELEMENTS if 0 < missing[e] < n]
        all_blank = [e for e in ELEMENTS if missing[e] == n]
        zero_and_blank = [
            e for e in ELEMENTS
            if missing[e] and any(finite_number(row.get(e)) and float(row[e]) == 0 for row in group)
        ]
        if not any(missing.values()):
            schema = "fully_populated"
        elif partial:
            schema = "mixed_within_source"
        else:
            schema = "consistent_source_schema_with_blanks"

        rows_all_blank = sum(all(row.get(e) is None for e in ELEMENTS) for row in group)
        rows_major_blank = sum(all(row.get(e) is None for e in MAJOR_ELEMENTS) for row in group)
        compositions = {normalized_tuple(row, ELEMENTS) for row in group}
        conditions = {normalized_tuple(row, CONDITION_FIELDS) for row in group}

        fe_remainders: list[float] = []
        fe_not_largest = 0
        for row in group:
            balance_elements = [e for e in ELEMENTS if not (source == "2000Russell" and e == "Ti")]
            recorded = [float(row[e]) for e in balance_elements if finite_number(row.get(e))]
            if not recorded:
                continue
            remainder = 100.0 - sum(recorded)
            fe_remainders.append(remainder)
            if remainder < max(recorded):
                fe_not_largest += 1

        if len(compositions) == 1 and len(conditions) == 1:
            row_structure = "repeated_measurements_single_composition_condition"
        elif len(compositions) == 1:
            row_structure = "single_composition_multiple_conditions"
        elif len(conditions) == 1:
            row_structure = "multiple_compositions_single_condition"
        else:
            row_structure = "multiple_compositions_and_conditions"

        title = meta.get("title", "")
        doi = clean(meta.get("doi", ""))
        if source == "2015Blackwood":
            title = "Can the Point Defect Model Explain the Influence of Temperature and Anion Size on Pitting of Stainless Steels"
            doi = "10.14773/cst.2015.14.6.253"

        base = {
            "material_family": "Fe-based stainless steel",
            "stainless_scope": "yes",
            "balance_element": "Fe_expected",
            "balance_reconstruction": "approximate_only_100_minus_recorded_wt_pct",
            "material_note": "Collection scope and source title/designation support an Fe-based stainless material; exact grade chemistry can contain unrecorded elements.",
            "evidence_level": "collection_scope_plus_source_title_or_designation",
            "evidence_url": f"https://doi.org/{doi}" if doi and doi.upper() != "NA" else (meta.get("oa_landing_url", "") or ""),
        }
        base.update(MATERIAL_OVERRIDES.get(source, {}))

        issue, correction = SOURCE_ISSUES.get(source, ("", ""))
        flags: list[str] = []
        if issue:
            flags.append(issue)
        if rows_all_blank:
            flags.append(f"{rows_all_blank}_rows_all_composition_blank")
        elif rows_major_blank:
            flags.append(f"{rows_major_blank}_rows_major_composition_blank")
        if zero_and_blank:
            flags.append("blank_and_explicit_zero_coexist:" + semicolon(zero_and_blank))
        if schema == "mixed_within_source":
            flags.append("composition_reporting_varies_within_source")

        if source in REPAIR_REQUIRED:
            review_status = "repair_required"
            handling = correction
        elif source in NON_FE:
            review_status = "keep_separate_non_Fe_family"
            handling = "Keep for a broad corrosion-resistant-alloy task or exclude from a stainless-only subset; never add Fe as balance."
        elif source in MIXED_FAMILY:
            review_status = "split_or_filter_material_families"
            handling = correction or "Split source rows by material family before reconstruction or filtering."
        elif source in BORDERLINE:
            review_status = "usable_with_material_taxonomy_caveat"
            handling = correction or "Keep for the comparison benchmark and retain grade/source identity; define the stainless-only taxonomy before a restricted model."
        elif issue or rows_all_blank == n or (rows_major_blank and rows_major_blank / n >= 0.5):
            review_status = "usable_with_composition_caveat"
            handling = correction or "Keep with source/grade indicator; recover composition before a chemistry-only final model."
        else:
            review_status = "usable_Fe_based"
            handling = "Keep for the comparative benchmark; preserve source grouping and apply the chosen missing-value policy consistently."

        event_counts = Counter(row.get("event") for row in group)
        numeric_target = sum(finite_number(row.get("E_pit")) for row in group)
        designations = sorted({str(row["alloy_designation"]).strip() for row in group if row.get("alloy_designation") not in (None, "")})

        records.append({
            "source": source,
            "source_id": meta.get("source_id", ""),
            "samples": n,
            "reference": meta.get("reference", ""),
            "doi": doi,
            "title": title,
            "alloy_designations": semicolon(designations),
            "material_family": base["material_family"],
            "stainless_scope": base["stainless_scope"],
            "balance_element": base["balance_element"],
            "balance_reconstruction": base["balance_reconstruction"],
            "rows_Fe_balance_not_largest_if_blanks_zero": fe_not_largest,
            "approx_Fe_balance_min_wt_pct": round(min(fe_remainders), 6) if fe_remainders else "",
            "approx_Fe_balance_max_wt_pct": round(max(fe_remainders), 6) if fe_remainders else "",
            "composition_basis": "wt_pct",
            "composition_provenance": "certificate_or_standard_at_collection_level; measured_vs_nominal_unresolved_per_source_unless_noted",
            "blank_semantics": "missing_or_unreported; not_verified_zero",
            "composition_schema": schema,
            "partial_elements": semicolon(partial),
            "all_blank_elements": semicolon(all_blank),
            "rows_all_composition_blank": rows_all_blank,
            "rows_major_composition_blank": rows_major_blank,
            "unique_compositions": len(compositions),
            "unique_test_conditions": len(conditions),
            "row_structure": row_structure,
            "numeric_breakdown_potential_rows": numeric_target,
            "event_1_pitting_rows": event_counts.get(1, 0),
            "event_0_censored_rows": event_counts.get(0, 0),
            "event_missing_rows": event_counts.get(None, 0),
            "target_convention": "E_pit_column_is_Ebr_mV_vs_AgAgCl_3M_KCl; may_be_pitting_or_competing_breakdown",
            "event_convention": "1=actual_pitting; 0=competing_non-pitting_breakdown_right-censored",
            "environment_convention": "concentrations_in_raw_workbook_are_M; paper_model_used_log10_and_1e-5_M_floor_for_unmentioned_anions",
            "source_specific_issue": semicolon(flags),
            "required_correction_or_caveat": correction,
            "material_note": base["material_note"],
            "evidence_level": base["evidence_level"],
            "evidence_url": base.get("evidence_url", ""),
            "collection_evidence": COLLECTION_PAPER,
            "classification_confidence": (
                "medium" if base["evidence_level"] in {
                    "workbook_composition_inference",
                    "secondary_bibliographic_identification_and_workbook_compositions",
                    "workbook_designations_and_compositions",
                } else "high"
            ),
            "review_status": review_status,
            "recommended_handling": handling,
        })

    fieldnames = list(records[0])
    with OUTPUT.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(records)

    status_sources = Counter(record["review_status"] for record in records)
    status_samples = Counter()
    scope_sources = Counter(record["stainless_scope"] for record in records)
    scope_samples = Counter()
    for record in records:
        status_samples[record["review_status"]] += int(record["samples"])
        scope_samples[record["stainless_scope"]] += int(record["samples"])

    summary = {
        "generated_from": str(WORKBOOK.relative_to(ROOT)),
        "sources": len(records),
        "samples": sum(int(record["samples"]) for record in records),
        "review_status_sources": dict(sorted(status_sources.items())),
        "review_status_samples": dict(sorted(status_samples.items())),
        "stainless_scope_sources": dict(sorted(scope_sources.items())),
        "stainless_scope_samples": dict(sorted(scope_samples.items())),
        "mixed_family_sources": sorted(MIXED_FAMILY),
        "non_Fe_sources": sorted(NON_FE),
        "repair_required_sources": sorted(REPAIR_REQUIRED),
        "global_conventions": {
            "composition": "wt%; certificate or standard; blanks missing/unreported",
            "target": "breakdown potential Ebr in mV vs Ag/AgCl (3 M KCl)",
            "event": "1 actual pitting; 0 competing breakdown/right censored",
            "raw_environment": "anion concentrations in M; log10/floor were paper-model preprocessing",
            "row_identity": "reported measurement/condition; rows within source are dependent",
        },
    }
    SUMMARY.write_text(json.dumps(summary, indent=2) + "\n", encoding="utf-8")

    assert len(records) == 154, len(records)
    assert sum(int(record["samples"]) for record in records) == 4460
    assert all(record["review_status"] for record in records)
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
